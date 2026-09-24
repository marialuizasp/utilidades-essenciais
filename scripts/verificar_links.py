"""Checagem conservadora de links HTTP e relatório diário."""
import os, ssl, smtplib, time
from datetime import datetime
from email.message import EmailMessage
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

ARQUIVO = "controle/controle_links_afiliados_utilidades_essenciais.xlsx"
AGORA = datetime.now(ZoneInfo("America/Sao_Paulo")).replace(tzinfo=None)

def verificar(url):
    if not isinstance(url, str) or not url.startswith(("https://", "http://")):
        return "Não verificado", "URL ausente ou inválida"
    req = Request(url, headers={"User-Agent": "Mozilla/5.0", "Accept": "text/html,application/xhtml+xml,*/*"})
    try:
        with urlopen(req, timeout=18) as resposta:
            destino = resposta.geturl()
            status = resposta.status
            if destino.rstrip("/") != url.rstrip("/"):
                return "Redirecionado", f"HTTP {status}; destino: {destino[:180]}"
            return "Funcionando", f"HTTP {status} (não confirma estoque nem preço)"
    except HTTPError as erro:
        if erro.code in (404, 410):
            return "Quebrado", f"HTTP {erro.code}: revisar manualmente"
        if erro.code in (401, 403, 429):
            return "Acesso bloqueado", f"HTTP {erro.code}: inconclusivo"
        return "Não verificado", f"HTTP {erro.code}"
    except (URLError, TimeoutError, OSError) as erro:
        return "Não verificado", f"Falha temporária: {str(erro)[:100]}"

def main():
    wb = load_workbook(ARQUIVO)
    aba = wb["Afiliados"]
    historico = wb["Historico"]
    resumo, mudancas = [], []
    for linha in range(2, aba.max_row + 1):
        produto = aba.cell(linha, 2).value
        if not produto:
            continue
        identificador = aba.cell(linha, 1).value or f"linha_{linha}"
        link = aba.cell(linha, 5).value
        anterior = aba.cell(linha, 11).value or "Não verificado"
        novo, detalhe = verificar(link)
        aba.cell(linha, 11).value = novo
        aba.cell(linha, 13).value = AGORA
        aba.cell(linha, 15).value = detalhe
        aba.cell(linha, 10).value = "Não verificado"
        if anterior != novo:
            historico.append([AGORA, identificador, produto, "Link", anterior, novo, "Checagem HTTP", link, detalhe])
            mudancas.append(f"- {produto}: {anterior} → {novo}")
        resumo.append(f"- {produto}: {novo}. {detalhe}")
        time.sleep(1)
    wb.save(ARQUIVO)
    texto = (
        f"Relatório diário — Utilidades Essenciais — {AGORA:%d/%m/%Y %H:%M}\n\n"
        f"Links consultados: {len(resumo)}\n\n"
        + "\n".join(resumo) + "\n\n"
        + "Mudanças:\n" + ("\n".join(mudancas) or "Nenhuma mudança detectada.")
        + "\n\nATENÇÃO: status HTTP não confirma anúncio disponível, estoque, preço ou comissão."
    )
    print(texto)
    required = ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "REPORT_TO"]
    if any(not os.getenv(k) for k in required):
        print("E-mail não enviado: falta configurar secrets SMTP e REPORT_TO.")
        return
    msg = EmailMessage()
    msg["Subject"] = f"Relatório de afiliados — {AGORA:%d/%m/%Y}"
    msg["From"] = os.environ["SMTP_USER"]
    msg["To"] = os.environ["REPORT_TO"]
    msg.set_content(texto)
    port = int(os.getenv("SMTP_PORT", "465"))
    with smtplib.SMTP_SSL(os.environ["SMTP_HOST"], port, context=ssl.create_default_context(), timeout=30) as smtp:
        smtp.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(msg)
    print("Relatório enviado.")

if __name__ == "__main__":
    main()
